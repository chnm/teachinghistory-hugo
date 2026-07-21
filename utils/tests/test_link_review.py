import importlib.util
from pathlib import Path

_spec = importlib.util.spec_from_file_location(
    "link_checker", Path(__file__).resolve().parent.parent / "link_checker.py"
)
lc = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(lc)


def test_normalize_heading_strips_punctuation_and_case():
    assert lc.normalize_heading("## For Further Reading") == "for further reading"
    assert lc.normalize_heading("References:") == "references"


def test_is_bibliography_heading_whole_match_only():
    assert lc.is_bibliography_heading("For Further Reading") is True
    assert lc.is_bibliography_heading("Bibliography") is True
    assert lc.is_bibliography_heading("References") is True
    # Must NOT match partial/compound headings:
    assert lc.is_bibliography_heading("Sources of Unity") is False
    assert lc.is_bibliography_heading("Sources and Teaching Models") is False


def test_heading_at_returns_nearest_preceding_heading():
    text = "# Intro\nbody\n## For Further Reading\n- a link here\n"
    idx = lc.build_heading_index(text)
    pos = text.index("a link here")
    assert lc.heading_at(idx, pos) == "For Further Reading"
    pos_intro = text.index("body")
    assert lc.heading_at(idx, pos_intro) == "Intro"


def test_looks_like_caption():
    assert lc.looks_like_caption("![poster](/img/x.jpg) A poster") is True
    assert lc.looks_like_caption("Courtesy of the Library of Congress, 1973") is True
    assert lc.looks_like_caption("Just a normal sentence with a point.") is False


def test_extract_links_with_context_classifies_kind_and_bucket_context():
    text = (
        "## For Further Reading\n"
        "See [the guide](https://example.org/guide) and also\n"
        "March to Wounded Knee, Library of Congress, https://www.loc.gov/item/2016648085/\n"
    )
    rows = lc.extract_links_with_context(text)
    by_url = {r["link_url"]: r for r in rows}

    guide = by_url["https://example.org/guide"]
    assert guide["link_kind"] == "hyperlink"
    assert guide["in_bibliography"] is True

    loc = by_url["https://www.loc.gov/item/2016648085/"]
    assert loc["link_kind"] == "bare_url"
    assert loc["in_bibliography"] is True
    assert loc["in_caption"] is True  # "Library of Congress" hint on the line


def test_is_reverifiable_status():
    assert lc.is_reverifiable_status("403") is True
    assert lc.is_reverifiable_status("404") is False
    assert lc.is_reverifiable_status("200") is False


class _StubResp:
    def __init__(self):
        self.status_code = 200
        self.url = "https://example.org/"
        self.headers = {"content-type": "text/html"}
        self.text = "<title>Fine</title>"


class _StubSession:
    def __init__(self):
        self.last_headers = None

    def get(self, url, timeout, allow_redirects, headers):
        self.last_headers = headers
        return _StubResp()


def test_check_url_uses_supplied_user_agent():
    session = _StubSession()
    lc.check_url("https://example.org/", session, user_agent=lc.BROWSER_UA)
    assert session.last_headers["User-Agent"] == lc.BROWSER_UA


class _RedirectStubResp:
    def __init__(self, final_url):
        self.status_code = 200
        self.url = final_url
        self.headers = {"content-type": "text/html"}
        self.text = "<title>Fine</title>"


class _RedirectStubSession:
    def __init__(self, final_url):
        self._final_url = final_url

    def get(self, url, timeout, allow_redirects, headers):
        return _RedirectStubResp(self._final_url)


def test_check_url_redirect_domain_change_detection():
    # A genuine domain change is flagged. Regression guard for the old
    # `.lstrip("www.")` bug: lstrip strips a char SET, so "web.com" -> "eb.com",
    # which made web.com -> eb.com look like the SAME domain. removeprefix fixes it.
    session = _RedirectStubSession("http://eb.com/page")
    result = lc.check_url("http://web.com/page", session)
    assert result["redirect_domain_changed"] is True

    # www normalization on the same host is NOT a domain change.
    session = _RedirectStubSession("http://example.org/")
    result = lc.check_url("http://www.example.org/", session)
    assert result["redirect_domain_changed"] is False


def test_classify_redirect():
    assert lc.classify_redirect("http://x.org/a", "https://x.org/a") == "benign"
    assert lc.classify_redirect("http://x.org/a", "http://www.x.org/a/") == "benign"
    assert lc.classify_redirect("http://x.org/a", "https://other.com/z") == "substantive"
    assert lc.classify_redirect("http://x.org/a", "") == "none"


def test_needs_wayback_lookup():
    # Dead statuses always qualify:
    assert lc.needs_wayback_lookup(
        {"link_url": "http://x.org/a", "http_status": "404", "final_url": ""}) is True
    # 200 with a substantive cross-host redirect (the C-candidate population):
    assert lc.needs_wayback_lookup(
        {"link_url": "http://x.org/a", "http_status": "200",
         "final_url": "https://other.com/z"}) is True
    # Same-host redirect is benign — no lookup:
    assert lc.needs_wayback_lookup(
        {"link_url": "http://x.org/a", "http_status": "200",
         "final_url": "https://www.x.org/a/"}) is False
    # Plain live link, no redirect:
    assert lc.needs_wayback_lookup(
        {"link_url": "http://x.org/a", "http_status": "200",
         "final_url": "http://x.org/a"}) is False
    # Never-checked row:
    assert lc.needs_wayback_lookup(
        {"link_url": "http://x.org/a", "http_status": "", "final_url": ""}) is False


def test_is_wayback_status_final():
    assert lc.is_wayback_status_final("found") is True
    assert lc.is_wayback_status_final("not_archived") is True
    # Blanks and transient failures get retried on --resume:
    assert lc.is_wayback_status_final("") is False
    assert lc.is_wayback_status_final("api_error_429") is False
    assert lc.is_wayback_status_final("api_error_503") is False
    assert lc.is_wayback_status_final(
        "error: HTTPSConnectionPool(host='archive.org', port=443): Read timed out.") is False


class _WaybackStubResp:
    def __init__(self, status_code, payload=None):
        self.status_code = status_code
        self._payload = payload or {}

    def json(self):
        return self._payload


class _RateLimitedThenOkSession:
    """First call returns 429, second returns a found snapshot."""
    def __init__(self):
        self.calls = 0

    def get(self, url, params, timeout, headers):
        self.calls += 1
        if self.calls == 1:
            return _WaybackStubResp(429)
        return _WaybackStubResp(200, {"archived_snapshots": {"closest": {
            "available": True, "url": "http://web.archive.org/web/1/x", "timestamp": "1"}}})


def test_lookup_wayback_retries_on_429(monkeypatch):
    sleeps = []
    monkeypatch.setattr(lc.time, "sleep", sleeps.append)
    session = _RateLimitedThenOkSession()
    result = lc.lookup_wayback("http://x.org/a", session)
    assert session.calls == 2
    assert result["wayback_status"] == "found"
    assert sleeps  # backed off before retrying


def test_lookup_wayback_gives_up_after_retries(monkeypatch):
    monkeypatch.setattr(lc.time, "sleep", lambda s: None)

    class _Always429:
        def get(self, url, params, timeout, headers):
            return _WaybackStubResp(429)

    result = lc.lookup_wayback("http://x.org/a", _Always429())
    assert result["wayback_status"] == "api_error_429"


def test_is_bookseller_and_bucket_priority():
    assert lc.is_bookseller("https://www.amazon.com/dp/123") is True
    assert lc.is_bookseller("https://books.google.com/books?id=1") is True
    assert lc.is_bookseller("https://loc.gov/item/1") is False
    # Bookseller wins even when under a bibliography heading:
    assert lc.bucket_for("https://amazon.com/x", True, False) == "bookseller"
    assert lc.bucket_for("https://loc.gov/x", True, False) == "bibliography"
    assert lc.bucket_for("https://loc.gov/x", False, True) == "caption_maybe"
    assert lc.bucket_for("https://loc.gov/x", False, False) == "none"


def test_broken_category():
    assert lc.broken_category("404", "none", "") == "A"
    assert lc.broken_category("CONN_ERROR", "none", "") == "A"
    assert lc.broken_category("403", "none", "") == "blocked-unknown"
    assert lc.broken_category("200", "substantive", "Padlet") == "C-candidate"
    assert lc.broken_category("200", "none", "404 Page Not Found") == "B?"
    assert lc.broken_category("200", "benign", "Real Article Title") == "live"
    assert lc.broken_category("429", "none", "") == "needs-human"
    assert lc.broken_category("", "none", "") == "needs-human"


def test_confidence_and_action():
    assert lc.confidence_for("A", "none") == "high"
    assert lc.confidence_for("C-candidate", "none") == "medium"
    assert lc.confidence_for("needs-human", "caption_maybe") == "low"

    assert lc.suggested_action("A", "bookseller", "not_archived") == "bulk-unlink"
    assert lc.suggested_action("A", "none", "found") == "salvage-wayback"
    assert lc.suggested_action("C-candidate", "none", "not_archived") == "needs-subjective-review"
    assert lc.suggested_action("live", "none", "") == "ok"


def test_write_master_xlsx(tmp_path):
    import openpyxl
    rows = [{k: "" for k in lc.MASTER_FIELDNAMES}]
    rows[0]["link_url"] = "https://example.org/x"
    rows[0]["bucket"] = "bookseller"
    out = tmp_path / "sheet.xlsx"
    lc.write_master_xlsx(rows, out)
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert [c.value for c in ws[1]] == lc.MASTER_FIELDNAMES + ["decision"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_is_homepage_url():
    assert lc.is_homepage_url("https://rrchnm.org") is True
    assert lc.is_homepage_url("https://rrchnm.org/") is True
    assert lc.is_homepage_url("https://example.org/index.html") is True
    assert lc.is_homepage_url("https://example.org/HOME/") is True
    assert lc.is_homepage_url("https://example.org/default.aspx") is True
    # Real paths and query strings are NOT homepages:
    assert lc.is_homepage_url("https://www.thirteen.org/wnet/historyofus/web07/segment2b.html") is False
    assert lc.is_homepage_url("https://example.org/?page_id=12") is False
    assert lc.is_homepage_url("https://www.c-span.org:443/series/?americanWriters") is False


def test_history_site_signal():
    # Title hit:
    assert "history" in lc.history_site_signal("A History of US", "https://x.com/a")
    # URL-only hit (row with no captured title):
    assert ".edu" in lc.history_site_signal("", "https://chnm.gmu.edu/page")
    # No hit:
    assert lc.history_site_signal("Casino online", "https://vn88.com/x") == []
    # Multiple hits all reported (sheet shows WHY something matched):
    hits = lc.history_site_signal("Museum of Education", "https://x.org/history/")
    assert set(hits) >= {"museum", "education", "history"}


def test_rebrand_verdict():
    # Homepage wins even when the title screams history (content gone, site survives):
    assert lc.rebrand_verdict("https://rrchnm.org", "Center for History and New Media") == "probably-broken"
    # Deep link + history signal:
    assert lc.rebrand_verdict(
        "https://www.thirteen.org/wnet/historyofus/web07.html",
        "Freedom: A History of US") == "probably-fine"
    # Deep link, no signal anywhere:
    assert lc.rebrand_verdict("https://vn88.com/casino", "Casino online") == "needs-human"
    # Missing title still matches on URL text:
    assert lc.rebrand_verdict("https://chnm.gmu.edu/loudountah/x", "") == "probably-fine"


def test_suggested_action_rebrand_verdicts():
    # probably-fine → mechanical repoint at the final URL:
    assert lc.suggested_action("C-candidate", "none", "",
                               rebrand_verdict="probably-fine") == "update-to-final-url"
    # probably-broken → wayback if archived, else remove/replace:
    assert lc.suggested_action("C-candidate", "none", "found",
                               rebrand_verdict="probably-broken") == "salvage-wayback"
    assert lc.suggested_action("C-candidate", "none", "not_archived",
                               rebrand_verdict="probably-broken") == "remove-or-replace"
    # needs-human keeps the old behavior:
    assert lc.suggested_action("C-candidate", "none", "",
                               rebrand_verdict="needs-human") == "needs-subjective-review"
    # Bucket still wins over any verdict:
    assert lc.suggested_action("C-candidate", "bookseller", "",
                               rebrand_verdict="probably-fine") == "bulk-unlink"


def _mk_link(source_file="history-content/website-reviews/foo-123.md",
             section="history-content", category="live", verdict="",
             wayback="", bulk=False):
    return {
        "section": section, "page_title": "Foo", "page_url": "/x/123",
        "source_file": source_file, "broken_category": category,
        "rebrand_verdict": verdict, "wayback_status": wayback,
        "bulk_delete_candidate": str(bulk),
    }


def test_subsection_of():
    assert lc.subsection_of("history-content/website-reviews/foo-1.md") == "website-reviews"
    assert lc.subsection_of("blog/a-post-2.md") == ""


def test_suggested_page_action():
    # Link-centric page with an un-archived broken link → delete candidate:
    assert lc.suggested_page_action("website-reviews", broken=1, broken_no_wayback=1) == "delete-page-candidate"
    # Link-centric, every broken link archived → salvage:
    assert lc.suggested_page_action("national-resources", broken=2, broken_no_wayback=0) == "salvage-wayback"
    # Link-centric but nothing broken → just fix links:
    assert lc.suggested_page_action("website-reviews", broken=0, broken_no_wayback=0) == "fix-links-only"
    # Content-centric sections never get delete suggestions:
    assert lc.suggested_page_action("teaching-guides", broken=5, broken_no_wayback=5) == "fix-links-only"


def test_aggregate_pages_counts_and_filtering():
    rows = [
        # Page 1: website review — one dead link without wayback, one live:
        _mk_link(category="A", wayback="not_archived"),
        _mk_link(category="live"),
        # Page 2: all live → excluded from the sheet:
        _mk_link(source_file="blog/fine-1.md", section="blog", category="live"),
        # Page 3: teaching guide — probably-broken rebrand that HAS a snapshot:
        _mk_link(source_file="teaching-materials/teaching-guides/g-9.md",
                 section="teaching-materials", category="C-candidate",
                 verdict="probably-broken", wayback="found", bulk=True),
    ]
    pages = {p["source_file"]: p for p in lc.aggregate_pages(rows)}

    assert "blog/fine-1.md" not in pages  # all-live pages excluded

    review = pages["history-content/website-reviews/foo-123.md"]
    assert review["subsection"] == "website-reviews"
    assert review["total_links"] == 2
    assert review["live"] == 1
    assert review["dead_A"] == 1
    assert review["broken_no_wayback"] == 1
    assert review["suggested_page_action"] == "delete-page-candidate"

    guide = pages["teaching-materials/teaching-guides/g-9.md"]
    assert guide["rebrand_probably_broken"] == 1
    assert guide["broken_with_wayback"] == 1
    assert guide["bulk_delete_candidates"] == 1
    assert guide["suggested_page_action"] == "fix-links-only"


def test_write_review_xlsx_generic(tmp_path):
    import openpyxl
    fieldnames = ["a", "b"]
    out = tmp_path / "pages.xlsx"
    lc.write_review_xlsx([{"a": 1, "b": "x"}], out, fieldnames, "page_review", {"a": 20})
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.title == "page_review"
    assert [c.value for c in ws[1]] == fieldnames
    assert ws.freeze_panes == "A2"


def test_write_review_xlsx_decision_dropdown(tmp_path):
    import openpyxl
    out = tmp_path / "sheet.xlsx"
    lc.write_review_xlsx([{"a": 1}], out, ["a"], "s", {},
                         decision_options=["remove", "keep"])
    ws = openpyxl.load_workbook(out).active
    # Decision header appended after the data columns, bold like the rest:
    assert ws.cell(row=1, column=2).value == "decision"
    assert ws.cell(row=1, column=2).font.bold
    # One list validation covering the decision column's data rows:
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    assert dvs[0].formula1 == '"remove,keep"'
    assert str(dvs[0].sqref) == "B2"  # openpyxl collapses the single-cell range
    # Autofilter extends over the decision column:
    assert ws.auto_filter.ref == "A1:B2"
