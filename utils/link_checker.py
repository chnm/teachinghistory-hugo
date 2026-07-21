# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "requests",
#     "beautifulsoup4",
#     "pyyaml",
#     "openpyxl",
# ]
# ///
"""
External link inventory and checker for teachinghistory.org Hugo site.

Usage:
    uv run utils/link_checker.py extract              # Build link inventory CSV
    uv run utils/link_checker.py check                # Check links from inventory
    uv run utils/link_checker.py check --resume       # Resume interrupted check run
    uv run utils/link_checker.py check --limit 100    # Check only first N unchecked links
    uv run utils/link_checker.py wayback              # Look up Wayback Machine snapshots for dead links
    uv run utils/link_checker.py wayback --resume     # Resume interrupted wayback run
    uv run utils/link_checker.py recheck              # Re-check 403s with a browser User-Agent
    uv run utils/link_checker.py recheck --limit 100  # Re-check only first N 403 URLs
    uv run utils/link_checker.py replace              # Dry run: show what would be replaced
    uv run utils/link_checker.py replace --apply      # Actually replace dead links with Wayback URLs
    uv run utils/link_checker.py classify             # Join inventory+results+wayback into master review sheet
    uv run utils/link_checker.py pages                # Aggregate master sheet into per-page review sheet
"""

import argparse
import collections
import csv
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import requests
import yaml
from bs4 import BeautifulSoup

CONTENT_DIR = Path(__file__).resolve().parent.parent / "teachinghistory-website" / "content"
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "utils"
INVENTORY_CSV = OUTPUT_DIR / "link_inventory.csv"
RESULTS_CSV = OUTPUT_DIR / "link_results.csv"

# Domains to exclude (internal links)
INTERNAL_DOMAINS = {
    "teachinghistory.org",
    "www.teachinghistory.org",
    "dev.teachinghistory.org",
    "localhost",
}

# Patterns in page titles that suggest a hijacked/parked domain
SUSPICIOUS_TITLE_WORDS = [
    "casino", "poker", "gambling", "slot", "betting",
    "buy this domain", "domain for sale", "parked", "is for sale",
    "adult", "xxx", "porn",
    "404", "not found", "page not found",
    "403", "forbidden", "access denied",
    "error", "default web page",
    "website expired", "account suspended", "account has been suspended",
    "coming soon", "under construction",
]

REQUEST_TIMEOUT = 15
REQUEST_DELAY = 0.3  # seconds between requests
USER_AGENT = "TeachingHistory-LinkChecker/1.0 (educational site maintenance)"
BROWSER_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/125.0.0.0 Safari/537.36"
)


# --- Extract phase ---

def parse_frontmatter(text: str) -> dict:
    """Extract YAML frontmatter from a markdown file."""
    match = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.DOTALL)
    if not match:
        return {}
    try:
        return yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError:
        return {}


BIBLIOGRAPHY_HEADINGS = {
    "for further reading", "further reading", "bibliography",
    "references", "works cited", "further resources",
    "related resources", "sources", "notes", "endnotes",
}

CAPTION_HINTS = re.compile(
    r"(library of congress|courtesy of|courtesy,|source:|photo:|image:|"
    r"credit:|national archives|smithsonian)",
    re.IGNORECASE,
)


def build_heading_index(text: str) -> list[tuple[int, str]]:
    """List of (char_offset, heading_text) for each ATX heading, in order.

    Only recognizes ATX headings (`#`-prefixed); setext headings (underlined
    with `===`/`---`) are not detected.
    """
    index = []
    offset = 0
    for line in text.splitlines(keepends=True):
        m = re.match(r"\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$", line)
        if m:
            index.append((offset, m.group(1).strip()))
        offset += len(line)
    return index


def heading_at(index: list[tuple[int, str]], pos: int) -> str:
    """Nearest heading at or before character offset `pos` ("" if none)."""
    current = ""
    for off, head in index:
        if off <= pos:
            current = head
        else:
            break
    return current


def normalize_heading(h: str) -> str:
    stripped = re.sub(r"^#+\s*", "", h).strip()
    return re.sub(r"[^a-z0-9 ]", "", stripped.lower()).strip()


def is_bibliography_heading(h: str) -> bool:
    return normalize_heading(h) in BIBLIOGRAPHY_HEADINGS


def looks_like_caption(line: str) -> bool:
    if "![" in line:
        return True
    return bool(CAPTION_HINTS.search(line))


def _line_at(text: str, pos: int) -> str:
    start = text.rfind("\n", 0, pos) + 1
    end = text.find("\n", pos)
    if end == -1:
        end = len(text)
    return text[start:end]


def extract_links_with_context(text: str) -> list[dict]:
    """Extract links with document context. Captures markdown links, HTML
    anchors, and written-out bare URLs (marked link_kind='bare_url')."""
    index = build_heading_index(text)
    spans: list[tuple[int, int]] = []
    results: list[dict] = []

    def add(pos: int, url: str, link_text: str, kind: str) -> None:
        heading = heading_at(index, pos)
        results.append({
            "link_url": url,
            "link_text": link_text,
            "link_kind": kind,
            "doc_heading": heading,
            "in_bibliography": is_bibliography_heading(heading),
            "in_caption": looks_like_caption(_line_at(text, pos)),
        })

    for m in re.finditer(r"\[([^\]]*)\]\(([^)]+)\)", text):
        spans.append((m.start(), m.end()))
        add(m.start(), m.group(2).strip(), m.group(1).strip(), "hyperlink")

    for m in re.finditer(
        r'<a\s[^>]*href=["\']([^"\']+)["\'][^>]*>(.*?)</a>',
        text, re.IGNORECASE | re.DOTALL,
    ):
        spans.append((m.start(), m.end()))
        link_text = re.sub(r"<[^>]+>", "", m.group(2)).strip()
        add(m.start(), m.group(1).strip(), link_text, "hyperlink")

    for m in re.finditer(r"https?://[^\s<>)\"'\]]+", text):
        if any(s <= m.start() < e for s, e in spans):
            continue  # already captured inside a markdown/HTML link
        url = m.group(0).rstrip(".,;")
        add(m.start(), url, url, "bare_url")

    return results


def is_external(url: str) -> bool:
    """Check if a URL is external (not teachinghistory.org or relative)."""
    if not url or url.startswith("#") or url.startswith("mailto:") or url.startswith("tel:"):
        return False
    parsed = urlparse(url)
    if not parsed.scheme and not parsed.netloc:
        return False  # relative link
    if parsed.scheme and parsed.scheme not in ("http", "https"):
        return False
    host = parsed.netloc.lower()
    return host not in INTERNAL_DOMAINS


def get_section(rel_path: Path) -> str:
    """Get the top-level section from a content file's relative path."""
    parts = rel_path.parts
    return parts[0] if parts else ""


def run_extract():
    """Scan all content files and build an external link inventory CSV."""
    print(f"Scanning {CONTENT_DIR} for external links...")
    rows = []
    file_count = 0
    for md_file in sorted(CONTENT_DIR.rglob("*.md")):
        file_count += 1
        text = md_file.read_text(encoding="utf-8", errors="replace")
        fm = parse_frontmatter(text)
        page_title = fm.get("title", "")
        page_url = fm.get("url", "")
        rel_path = md_file.relative_to(CONTENT_DIR)
        section = get_section(rel_path)

        for link in extract_links_with_context(text):
            url = link["link_url"]
            if is_external(url):
                rows.append({
                    "section": section,
                    "page_title": page_title,
                    "page_url": page_url,
                    "source_file": str(rel_path),
                    "link_url": url,
                    "link_text": link["link_text"],
                    "link_kind": link["link_kind"],
                    "doc_heading": link["doc_heading"],
                    "in_bibliography": link["in_bibliography"],
                    "in_caption": link["in_caption"],
                })

    # Dedupe exact (source_file, link_url) pairs
    seen = set()
    deduped = []
    for row in rows:
        key = (row["source_file"], row["link_url"])
        if key not in seen:
            seen.add(key)
            deduped.append(row)

    fieldnames = [
        "section", "page_title", "page_url", "source_file",
        "link_url", "link_text", "link_kind",
        "doc_heading", "in_bibliography", "in_caption",
    ]
    with open(INVENTORY_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(deduped)

    unique_urls = len({r["link_url"] for r in deduped})
    print(f"Scanned {file_count} files.")
    print(f"Found {len(deduped)} external link references ({unique_urls} unique URLs).")
    print(f"Inventory written to {INVENTORY_CSV}")


# --- Check phase ---

def check_url(url: str, session: requests.Session, user_agent: str = USER_AGENT) -> dict:
    """Check a single URL. Returns status info dict."""
    result = {
        "http_status": "",
        "final_url": "",
        "remote_title": "",
        "redirect_domain_changed": False,
        "needs_review": False,
        "reason": "",
    }
    try:
        resp = session.get(
            url,
            timeout=REQUEST_TIMEOUT,
            allow_redirects=True,
            headers={"User-Agent": user_agent},
        )
        result["http_status"] = resp.status_code
        result["final_url"] = resp.url

        # Check if redirect changed domain
        orig_domain = urlparse(url).netloc.lower().removeprefix("www.")
        final_domain = urlparse(resp.url).netloc.lower().removeprefix("www.")
        if orig_domain != final_domain:
            result["redirect_domain_changed"] = True

        # Extract title for 200 responses
        if resp.status_code == 200 and "text/html" in resp.headers.get("content-type", ""):
            soup = BeautifulSoup(resp.text[:50000], "html.parser")
            title_tag = soup.find("title")
            if title_tag:
                result["remote_title"] = title_tag.get_text(strip=True)[:200]

        # Determine if needs review
        reasons = []
        if resp.status_code >= 400:
            reasons.append(f"HTTP {resp.status_code}")
        if result["redirect_domain_changed"]:
            reasons.append(f"redirected to different domain: {final_domain}")
        title_lower = result["remote_title"].lower()
        for word in SUSPICIOUS_TITLE_WORDS:
            if word in title_lower:
                reasons.append(f"suspicious title keyword: '{word}'")
                break

        if reasons:
            result["needs_review"] = True
            result["reason"] = "; ".join(reasons)

    except requests.exceptions.SSLError:
        result["http_status"] = "SSL_ERROR"
        result["needs_review"] = True
        result["reason"] = "SSL certificate error"
    except requests.exceptions.ConnectionError:
        result["http_status"] = "CONN_ERROR"
        result["needs_review"] = True
        result["reason"] = "Connection failed (site may be down)"
    except requests.exceptions.Timeout:
        result["http_status"] = "TIMEOUT"
        result["needs_review"] = True
        result["reason"] = f"Timed out after {REQUEST_TIMEOUT}s"
    except requests.exceptions.TooManyRedirects:
        result["http_status"] = "TOO_MANY_REDIRECTS"
        result["needs_review"] = True
        result["reason"] = "Too many redirects"
    except requests.exceptions.RequestException as e:
        result["http_status"] = "ERROR"
        result["needs_review"] = True
        result["reason"] = str(e)[:200]

    return result


def run_check(resume: bool = False, limit: int | None = None):
    """Read inventory CSV, check each unique URL, write results CSV."""
    if not INVENTORY_CSV.exists():
        print(f"No inventory found at {INVENTORY_CSV}. Run 'extract' first.")
        sys.exit(1)

    # Read inventory
    with open(INVENTORY_CSV, encoding="utf-8") as f:
        inventory = list(csv.DictReader(f))
    print(f"Loaded {len(inventory)} link references from inventory.")

    # Get unique URLs
    unique_urls = sorted({row["link_url"] for row in inventory})
    print(f"Found {len(unique_urls)} unique URLs to check.")

    # Load already-checked URLs if resuming (only rows with actual results)
    checked = {}
    if resume and RESULTS_CSV.exists():
        with open(RESULTS_CSV, encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if row.get("http_status", ""):
                    checked[row["link_url"]] = row
        print(f"Resuming: {len(checked)} URLs already checked.")

    # Check URLs
    url_results = {}
    urls_to_check = [u for u in unique_urls if u not in checked]
    if limit:
        urls_to_check = urls_to_check[:limit]
    total = len(urls_to_check)
    print(f"Checking {total} URLs...")

    session = requests.Session()
    for i, url in enumerate(urls_to_check, 1):
        if i % 50 == 0 or i == 1:
            print(f"  [{i}/{total}] checking {url[:80]}...")
        result = check_url(url, session)
        url_results[url] = result
        time.sleep(REQUEST_DELAY)

    # Merge with previously checked
    for url, row in checked.items():
        url_results[url] = {
            "http_status": row.get("http_status", ""),
            "final_url": row.get("final_url", ""),
            "remote_title": row.get("remote_title", ""),
            "redirect_domain_changed": row.get("redirect_domain_changed", ""),
            "needs_review": row.get("needs_review", ""),
            "reason": row.get("reason", ""),
        }

    # Build output: one row per (source_file, link_url) with check results merged
    fieldnames = [
        "section", "page_title", "page_url", "source_file",
        "link_url", "link_text",
        "http_status", "final_url", "remote_title",
        "redirect_domain_changed", "needs_review", "reason",
    ]
    output_rows = []
    for row in inventory:
        url = row["link_url"]
        result = url_results.get(url, {})
        output_rows.append({
            **row,
            "http_status": result.get("http_status", ""),
            "final_url": result.get("final_url", ""),
            "remote_title": result.get("remote_title", ""),
            "redirect_domain_changed": result.get("redirect_domain_changed", ""),
            "needs_review": result.get("needs_review", ""),
            "reason": result.get("reason", ""),
        })

    with open(RESULTS_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(output_rows)

    flagged = sum(1 for r in output_rows if str(r["needs_review"]) == "True")
    print(f"\nResults written to {RESULTS_CSV}")
    print(f"Total link references: {len(output_rows)}")
    print(f"Flagged for review: {flagged}")


def is_reverifiable_status(status: str) -> bool:
    """Statuses worth re-checking with a browser UA (bot-block false positives)."""
    return str(status) == "403"


def run_recheck(limit: int | None = None):
    """Re-check 403 URLs in RESULTS_CSV with a browser-like User-Agent."""
    if not RESULTS_CSV.exists():
        print(f"No results found at {RESULTS_CSV}. Run 'check' first.")
        sys.exit(1)

    with open(RESULTS_CSV, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = list(reader.fieldnames)

    target_urls = sorted({
        r["link_url"] for r in rows if is_reverifiable_status(r.get("http_status", ""))
    })
    if limit:
        target_urls = target_urls[:limit]
    print(f"Re-checking {len(target_urls)} URLs (403) with browser User-Agent...")

    session = requests.Session()
    updated: dict[str, dict] = {}
    for i, url in enumerate(target_urls, 1):
        if i % 50 == 0 or i == 1:
            print(f"  [{i}/{len(target_urls)}] rechecking {url[:80]}...")
        updated[url] = check_url(url, session, user_agent=BROWSER_UA)
        time.sleep(REQUEST_DELAY)

    changed = 0
    for r in rows:
        u = r["link_url"]
        if u in updated:
            res = updated[u]
            if str(r["http_status"]) != str(res["http_status"]):
                changed += 1
            r["http_status"] = res["http_status"]
            r["final_url"] = res["final_url"]
            r["remote_title"] = res["remote_title"]
            r["redirect_domain_changed"] = res["redirect_domain_changed"]
            r["needs_review"] = res["needs_review"]
            r["reason"] = res["reason"]

    with open(RESULTS_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    print(f"\nRewrote {RESULTS_CSV}. {changed} rows changed status after re-check.")


# --- Wayback phase ---

WAYBACK_CSV = OUTPUT_DIR / "link_wayback.csv"
WAYBACK_API = "https://archive.org/wayback/available"
WAYBACK_DELAY = 0.5  # be polite to archive.org

# Statuses that indicate a dead/broken link worth looking up in Wayback
DEAD_STATUSES = {
    "404", "410", "CONN_ERROR", "TIMEOUT", "SSL_ERROR",
    "ERROR", "TOO_MANY_REDIRECTS", "500", "502", "503", "504",
    "520", "521", "530", "526",
}


WAYBACK_429_BACKOFF = 60  # seconds to wait before retrying a rate-limited lookup
WAYBACK_429_RETRIES = 2


def lookup_wayback(url: str, session: requests.Session) -> dict:
    """Query the Wayback Machine Availability API for a snapshot."""
    result = {"wayback_url": "", "wayback_timestamp": "", "wayback_status": ""}
    for attempt in range(WAYBACK_429_RETRIES + 1):
        try:
            resp = session.get(
                WAYBACK_API,
                params={"url": url},
                timeout=REQUEST_TIMEOUT,
                headers={"User-Agent": USER_AGENT},
            )
            if resp.status_code == 429 and attempt < WAYBACK_429_RETRIES:
                time.sleep(WAYBACK_429_BACKOFF)
                continue
            if resp.status_code == 200:
                data = resp.json()
                snapshot = data.get("archived_snapshots", {}).get("closest", {})
                if snapshot and snapshot.get("available"):
                    result["wayback_url"] = snapshot.get("url", "")
                    result["wayback_timestamp"] = snapshot.get("timestamp", "")
                    result["wayback_status"] = "found"
                else:
                    result["wayback_status"] = "not_archived"
            else:
                result["wayback_status"] = f"api_error_{resp.status_code}"
        except requests.exceptions.RequestException as e:
            result["wayback_status"] = f"error: {str(e)[:100]}"
        break
    return result


def needs_wayback_lookup(result_row: dict) -> bool:
    """A results row worth a Wayback lookup: dead, or a cross-host redirect
    (the C-candidate population, whose original URL may still be archived)."""
    if result_row.get("http_status", "") in DEAD_STATUSES:
        return True
    return classify_redirect(
        result_row.get("link_url", ""), result_row.get("final_url", "")) == "substantive"


def is_wayback_status_final(status: str) -> bool:
    """Statuses --resume keeps; blanks and transient failures get retried."""
    if not status:
        return False
    return not (status.startswith("api_error") or status.startswith("error:"))


def run_wayback(resume: bool = False, limit: int | None = None):
    """Look up Wayback Machine snapshots for dead and rebrand-suspect links.

    Processes two sources of URLs:
    1. URLs from link_results.csv that are dead (DEAD_STATUSES) or redirect
       cross-host (the C-candidate population) — see needs_wayback_lookup
    2. Any URLs already in link_wayback.csv without a final wayback_status
       (blank, api_error_*, or error: rows get looked up again)
    """
    if not RESULTS_CSV.exists():
        print(f"No results found at {RESULTS_CSV}. Run 'check' first.")
        sys.exit(1)

    # Read full results
    with open(RESULTS_CSV, encoding="utf-8") as f:
        results = list(csv.DictReader(f))

    # Get unique dead/cross-host-redirect URLs from results
    target_urls = {r["link_url"] for r in results if needs_wayback_lookup(r)}

    # Also include any URLs in the wayback CSV without a final status yet
    # (e.g. manually added flagged 200s, or earlier rate-limited lookups)
    wayback_rows = []
    if WAYBACK_CSV.exists():
        with open(WAYBACK_CSV, encoding="utf-8") as f:
            wayback_rows = list(csv.DictReader(f))
        for row in wayback_rows:
            if not is_wayback_status_final(row.get("wayback_status", "")):
                target_urls.add(row["link_url"])

    target_urls = sorted(target_urls)
    print(f"Found {len(target_urls)} unique URLs to look up in Wayback Machine.")

    # Load already-looked-up URLs if resuming
    checked = {}
    if resume:
        for row in wayback_rows:
            if is_wayback_status_final(row.get("wayback_status", "")):
                checked[row["link_url"]] = row
        print(f"Resuming: {len(checked)} URLs already looked up.")

    urls_to_check = [u for u in target_urls if u not in checked]
    if limit:
        urls_to_check = urls_to_check[:limit]
    total = len(urls_to_check)
    print(f"Looking up {total} URLs...")

    wayback_results = {}
    session = requests.Session()
    for i, url in enumerate(urls_to_check, 1):
        if i % 50 == 0 or i == 1:
            print(f"  [{i}/{total}] looking up {url[:80]}...")
        wayback_results[url] = lookup_wayback(url, session)
        time.sleep(WAYBACK_DELAY)

    # Merge with previously checked
    for url, row in checked.items():
        wayback_results[url] = {
            "wayback_url": row.get("wayback_url", ""),
            "wayback_timestamp": row.get("wayback_timestamp", ""),
            "wayback_status": row.get("wayback_status", ""),
        }

    fieldnames = [
        "section", "page_title", "page_url", "source_file",
        "link_url", "link_text",
        "http_status", "reason",
        "wayback_url", "wayback_timestamp", "wayback_status",
    ]

    output_rows = []
    seen_keys = set()

    # First: rows from results CSV needing lookups (dead or cross-host redirect)
    for r in results:
        if not needs_wayback_lookup(r):
            continue
        url = r["link_url"]
        key = (r["source_file"], url)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        wb = wayback_results.get(url, {})
        output_rows.append({
            "section": r["section"],
            "page_title": r["page_title"],
            "page_url": r["page_url"],
            "source_file": r["source_file"],
            "link_url": url,
            "link_text": r["link_text"],
            "http_status": r["http_status"],
            "reason": r["reason"],
            "wayback_url": wb.get("wayback_url", ""),
            "wayback_timestamp": wb.get("wayback_timestamp", ""),
            "wayback_status": wb.get("wayback_status", ""),
        })

    # Second: remaining rows from the existing wayback CSV
    # (e.g. flagged 200s added manually)
    for row in wayback_rows:
        key = (row["source_file"], row["link_url"])
        if key in seen_keys:
            continue
        seen_keys.add(key)
        url = row["link_url"]
        wb = wayback_results.get(url, {})
        output_rows.append({
            "section": row["section"],
            "page_title": row["page_title"],
            "page_url": row["page_url"],
            "source_file": row["source_file"],
            "link_url": url,
            "link_text": row["link_text"],
            "http_status": row["http_status"],
            "reason": row["reason"],
            "wayback_url": wb.get("wayback_url", ""),
            "wayback_timestamp": wb.get("wayback_timestamp", ""),
            "wayback_status": wb.get("wayback_status", ""),
        })

    with open(WAYBACK_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(output_rows)

    found = sum(1 for r in output_rows if r["wayback_status"] == "found")
    not_archived = sum(1 for r in output_rows if r["wayback_status"] == "not_archived")
    print(f"\nResults written to {WAYBACK_CSV}")
    print(f"Total link references: {len(output_rows)}")
    print(f"Wayback snapshot found:     {found}")
    print(f"Not archived:               {not_archived}")


# --- Replace phase ---

def run_replace(dry_run: bool = True):
    """Replace dead links in content files with Wayback Machine URLs."""
    if not WAYBACK_CSV.exists():
        print(f"No wayback results found at {WAYBACK_CSV}. Run 'wayback' first.")
        sys.exit(1)

    # Build replacement map: {(source_file, old_url): wayback_url}
    replacements = {}
    with open(WAYBACK_CSV, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["wayback_status"] == "found" and row["wayback_url"]:
                key = (row["source_file"], row["link_url"])
                replacements[key] = row["wayback_url"]

    # Group by source file
    by_file: dict[str, list[tuple[str, str]]] = {}
    for (source_file, old_url), new_url in replacements.items():
        by_file.setdefault(source_file, []).append((old_url, new_url))

    print(f"Found {len(replacements)} link replacements across {len(by_file)} files.")
    if dry_run:
        print("DRY RUN — no files will be modified. Use --apply to make changes.\n")

    files_modified = 0
    links_replaced = 0
    skipped = []

    for source_file, swaps in sorted(by_file.items()):
        file_path = CONTENT_DIR / source_file
        if not file_path.exists():
            skipped.append((source_file, "file not found"))
            continue

        text = file_path.read_text(encoding="utf-8")
        original = text
        file_replacements = 0

        for old_url, new_url in swaps:
            if old_url in text:
                text = text.replace(old_url, new_url)
                file_replacements += 1
            else:
                skipped.append((source_file, f"URL not found in file: {old_url[:80]}"))

        if text != original:
            if dry_run:
                print(f"  Would modify {source_file} ({file_replacements} links)")
            else:
                file_path.write_text(text, encoding="utf-8")
                print(f"  Modified {source_file} ({file_replacements} links)")
            files_modified += 1
            links_replaced += file_replacements

    print(f"\n{'Would replace' if dry_run else 'Replaced'} {links_replaced} links in {files_modified} files.")
    if skipped:
        print(f"Skipped {len(skipped)} replacements:")
        for source_file, reason in skipped[:20]:
            print(f"  {source_file}: {reason}")
        if len(skipped) > 20:
            print(f"  ... and {len(skipped) - 20} more")
    if dry_run:
        print("\nRun with --apply to make changes.")


# --- Classify phase ---

MASTER_CSV = OUTPUT_DIR / "link_review_master.csv"
MASTER_XLSX = OUTPUT_DIR / "link_review_master.xlsx"
PAGES_CSV = OUTPUT_DIR / "link_review_pages.csv"
PAGES_XLSX = OUTPUT_DIR / "link_review_pages.xlsx"
PAGES_XLSX_WIDTHS = {"page_title": 40, "source_file": 50, "page_url": 30}

BOOKSELLER_DOMAINS = {
    "amazon.", "barnesandnoble.", "bn.com", "abebooks.",
    "powells.", "bookshop.org", "thriftbooks.", "alibris.",
    "books.google.",
}


def _host(url: str) -> str:
    return urlparse(url).netloc.lower().removeprefix("www.")


def classify_redirect(orig_url: str, final_url: str) -> str:
    """'none' if no redirect, 'benign' if same host, 'substantive' if host changed."""
    if not final_url or final_url == orig_url:
        return "none"
    if _host(orig_url) == _host(final_url):
        return "benign"
    return "substantive"


HOMEPAGE_PATHS = {"", "index.html", "index.htm", "index.php", "home", "default.aspx"}

HISTORY_VOCAB = [
    "history", "historical", "primary source", "museum", "archive",
    "education", "teaching", "teacher", "library", "heritage",
    "humanities", "smithsonian", "social studies", "civics", "k-12",
    "learning", "lesson", "curriculum", "university", ".edu", ".gov",
]


def is_homepage_url(url: str) -> bool:
    """True when the URL points at a site root (trivial path, no query)."""
    parsed = urlparse(url)
    if parsed.query:
        return False
    return parsed.path.strip("/").lower() in HOMEPAGE_PATHS


def history_site_signal(remote_title: str, final_url: str) -> list[str]:
    """History/education vocabulary matched in the page title or final URL."""
    haystack = f"{remote_title or ''} {final_url or ''}".lower()
    return [word for word in HISTORY_VOCAB if word in haystack]


def rebrand_verdict(final_url: str, remote_title: str) -> str:
    """Triage a C-candidate: homepage redirect beats any history signal."""
    if is_homepage_url(final_url):
        return "probably-broken"
    if history_site_signal(remote_title, final_url):
        return "probably-fine"
    return "needs-human"


def is_bookseller(url: str) -> bool:
    host = _host(url)
    return any(token in host for token in BOOKSELLER_DOMAINS)


def bucket_for(url: str, in_bibliography: bool, in_caption: bool) -> str:
    if is_bookseller(url):
        return "bookseller"
    if in_bibliography:
        return "bibliography"
    if in_caption:
        return "caption_maybe"
    return "none"


def broken_category(http_status: str, redirect_kind: str, remote_title: str) -> str:
    status = str(http_status)
    if status in DEAD_STATUSES:
        return "A"
    if status == "403":
        return "blocked-unknown"
    if status.startswith("2"):
        title = (remote_title or "").lower()
        if any(w in title for w in SUSPICIOUS_TITLE_WORDS):
            return "B?"
        if redirect_kind == "substantive":
            return "C-candidate"
        return "live"
    return "needs-human"


def confidence_for(category: str, bucket: str) -> str:
    if category == "A" or bucket == "bookseller" or category == "live":
        return "high"
    if category in ("C-candidate", "B?") or bucket == "bibliography":
        return "medium"
    return "low"


def suggested_action(category: str, bucket: str, wayback_status: str,
                     rebrand_verdict: str = "") -> str:
    if bucket != "none":
        return "bulk-unlink"
    if category == "C-candidate":
        if rebrand_verdict == "probably-fine":
            return "update-to-final-url"
        if rebrand_verdict == "probably-broken":
            return "salvage-wayback" if wayback_status == "found" else "remove-or-replace"
        return "needs-subjective-review"
    if category == "A" and wayback_status == "found":
        return "salvage-wayback"
    if category in ("A", "B?", "blocked-unknown"):
        return "needs-subjective-review"
    return "ok"


MASTER_FIELDNAMES = [
    "section", "page_title", "page_url", "source_file",
    "link_url", "link_text", "link_kind",
    "doc_heading", "in_bibliography", "in_caption",
    "http_status", "final_url", "redirect_kind", "remote_title",
    "broken_category", "bucket", "bulk_delete_candidate",
    "confidence", "suggested_action",
    "wayback_url", "wayback_status",
    "final_is_homepage", "history_signal", "rebrand_verdict",
]


def write_review_xlsx(rows: list[dict], path: Path, fieldnames: list[str],
                      sheet_title: str, widths: dict[str, int],
                      decision_options: list[str] | None = None) -> None:
    """Write a review sheet as .xlsx with a frozen header and autofilter.

    decision_options adds a blank `decision` column with a dropdown so the
    team can record their call on each row.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(fieldnames)
    for row in rows:
        ws.append([row.get(col, "") for col in fieldnames])

    n_cols = len(fieldnames)
    if decision_options:
        n_cols += 1
        letter = ws.cell(row=1, column=n_cols).column_letter
        ws.cell(row=1, column=n_cols, value="decision")
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(decision_options) + '"',
                            allow_blank=True)
        dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")
        ws.add_data_validation(dv)
        ws.column_dimensions[letter].width = 22

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=n_cols).column_letter}{ws.max_row}"
    for i, col in enumerate(fieldnames, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(col, 16)
    wb.save(path)


MASTER_XLSX_WIDTHS = {"page_title": 40, "source_file": 40, "link_url": 50,
                      "link_text": 30, "final_url": 40, "remote_title": 30,
                      "history_signal": 30}

# Dropdown choices for the team's decision column (per link / per page).
LINK_DECISIONS = ["remove", "update-to-final-url", "salvage-wayback",
                  "salvage-new-link", "keep"]
PAGE_DECISIONS = ["delete-page", "salvage-wayback", "fix-links", "keep"]


def write_master_xlsx(rows: list[dict], path: Path) -> None:
    write_review_xlsx(rows, path, MASTER_FIELDNAMES, "link_review", MASTER_XLSX_WIDTHS,
                      decision_options=LINK_DECISIONS)


LINK_CENTRIC_SUBSECTIONS = {"website-reviews", "national-resources"}

PAGES_FIELDNAMES = [
    "section", "subsection", "page_title", "page_url", "source_file",
    "total_links", "live", "dead_A",
    "rebrand_probably_broken", "rebrand_probably_fine", "rebrand_needs_human",
    "soft404_B", "blocked_unknown", "needs_human", "bulk_delete_candidates",
    "broken_with_wayback", "broken_no_wayback",
    "suggested_page_action",
]


def subsection_of(source_file: str) -> str:
    """Second path segment (section/subsection/file.md), or '' when flat."""
    parts = Path(source_file).parts
    return parts[1] if len(parts) > 2 else ""


def suggested_page_action(subsection: str, broken: int, broken_no_wayback: int) -> str:
    """Advisory page-level verdict. 'Broken' = dead_A + rebrand_probably_broken."""
    if subsection in LINK_CENTRIC_SUBSECTIONS:
        if broken_no_wayback >= 1:
            return "delete-page-candidate"
        if broken >= 1:
            return "salvage-wayback"
    return "fix-links-only"


def aggregate_pages(rows: list[dict]) -> list[dict]:
    """One row per page with >=1 non-live link. Accepts CSV rows (str values)."""
    by_page: dict[str, list[dict]] = {}
    for r in rows:
        by_page.setdefault(r["source_file"], []).append(r)

    out = []
    for source_file, links in sorted(by_page.items()):
        cats = collections.Counter(r["broken_category"] for r in links)
        if cats.get("live", 0) == len(links):
            continue
        verdicts = collections.Counter(
            r.get("rebrand_verdict", "") for r in links if r.get("rebrand_verdict"))
        broken = [r for r in links
                  if r["broken_category"] == "A"
                  or r.get("rebrand_verdict") == "probably-broken"]
        with_wayback = sum(1 for r in broken if r.get("wayback_status") == "found")
        subsection = subsection_of(source_file)
        first = links[0]
        out.append({
            "section": first["section"],
            "subsection": subsection,
            "page_title": first["page_title"],
            "page_url": first["page_url"],
            "source_file": source_file,
            "total_links": len(links),
            "live": cats.get("live", 0),
            "dead_A": cats.get("A", 0),
            "rebrand_probably_broken": verdicts.get("probably-broken", 0),
            "rebrand_probably_fine": verdicts.get("probably-fine", 0),
            "rebrand_needs_human": verdicts.get("needs-human", 0),
            "soft404_B": cats.get("B?", 0),
            "blocked_unknown": cats.get("blocked-unknown", 0),
            "needs_human": cats.get("needs-human", 0),
            "bulk_delete_candidates": sum(
                1 for r in links if str(r.get("bulk_delete_candidate")) == "True"),
            "broken_with_wayback": with_wayback,
            "broken_no_wayback": len(broken) - with_wayback,
            "suggested_page_action": suggested_page_action(
                subsection, len(broken), len(broken) - with_wayback),
        })
    return out


def _load_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def run_classify():
    """Join inventory + results + wayback into the master review sheet."""
    inventory = _load_csv(INVENTORY_CSV)
    if not inventory:
        print(f"No inventory found at {INVENTORY_CSV}. Run 'extract' first.")
        sys.exit(1)

    results = {(r["source_file"], r["link_url"]): r for r in _load_csv(RESULTS_CSV)}
    wayback = {(r["source_file"], r["link_url"]): r for r in _load_csv(WAYBACK_CSV)}

    rows = []
    for inv in inventory:
        key = (inv["source_file"], inv["link_url"])
        res = results.get(key, {})
        wb = wayback.get(key, {})

        http_status = res.get("http_status", "")
        final_url = res.get("final_url", "")
        remote_title = res.get("remote_title", "")
        in_bib = str(inv.get("in_bibliography", "")) == "True"
        in_cap = str(inv.get("in_caption", "")) == "True"

        redirect_kind = classify_redirect(inv["link_url"], final_url)
        bucket = bucket_for(inv["link_url"], in_bib, in_cap)
        category = broken_category(http_status, redirect_kind, remote_title)
        wb_status = wb.get("wayback_status", "")

        if category == "C-candidate":
            is_home = is_homepage_url(final_url)
            signal = history_site_signal(remote_title, final_url)
            verdict = rebrand_verdict(final_url, remote_title)
        else:
            is_home, signal, verdict = "", [], ""

        rows.append({
            "section": inv["section"],
            "page_title": inv["page_title"],
            "page_url": inv["page_url"],
            "source_file": inv["source_file"],
            "link_url": inv["link_url"],
            "link_text": inv["link_text"],
            "link_kind": inv.get("link_kind", ""),
            "doc_heading": inv.get("doc_heading", ""),
            "in_bibliography": in_bib,
            "in_caption": in_cap,
            "http_status": http_status,
            "final_url": final_url,
            "redirect_kind": redirect_kind,
            "remote_title": remote_title,
            "broken_category": category,
            "bucket": bucket,
            "bulk_delete_candidate": bucket != "none",
            "confidence": confidence_for(category, bucket),
            "suggested_action": suggested_action(category, bucket, wb_status, verdict),
            "wayback_url": wb.get("wayback_url", ""),
            "wayback_status": wb_status,
            "final_is_homepage": is_home,
            "history_signal": "; ".join(signal),
            "rebrand_verdict": verdict,
        })

    with open(MASTER_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=MASTER_FIELDNAMES, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    n_candidates = sum(1 for r in rows if r["bulk_delete_candidate"])
    print(f"\nWrote {MASTER_CSV} ({len(rows)} rows).")
    print(f"Bulk-delete candidates: {n_candidates}")
    cats = collections.Counter(r["broken_category"] for r in rows)
    for cat, n in cats.most_common():
        print(f"  {cat}: {n}")

    verdicts = collections.Counter(r["rebrand_verdict"] for r in rows if r["rebrand_verdict"])
    if verdicts:
        print("Rebrand verdicts (C-candidate):")
        for v, n in verdicts.most_common():
            print(f"  {v}: {n}")

    write_master_xlsx(rows, MASTER_XLSX)
    print(f"Wrote {MASTER_XLSX}")

    unchecked = sum(1 for r in rows if not r["http_status"])
    if unchecked:
        print(f"\nWARNING: {unchecked} links have no HTTP status (never checked) — "
              f"they are categorized 'needs-human' until you run `check` to cover them. "
              f"Filter http_status == '' in the sheet to find them.")


def run_pages():
    """Aggregate the master sheet into one row per page with problem links."""
    rows = _load_csv(MASTER_CSV)
    if not rows:
        print(f"No master sheet at {MASTER_CSV}. Run 'classify' first.")
        sys.exit(1)
    if "rebrand_verdict" not in rows[0]:
        print(f"{MASTER_CSV} predates the rebrand-verdict columns. Re-run 'classify' first.")
        sys.exit(1)

    pages = aggregate_pages(rows)

    with open(PAGES_CSV, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PAGES_FIELDNAMES, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(pages)
    print(f"Wrote {PAGES_CSV} ({len(pages)} pages with problem links).")

    write_review_xlsx(pages, PAGES_XLSX, PAGES_FIELDNAMES, "page_review", PAGES_XLSX_WIDTHS,
                      decision_options=PAGE_DECISIONS)
    print(f"Wrote {PAGES_XLSX}")

    actions = collections.Counter(p["suggested_page_action"] for p in pages)
    for action, n in actions.most_common():
        print(f"  {action}: {n}")


# --- CLI ---

def main():
    parser = argparse.ArgumentParser(description="External link checker for teachinghistory.org")
    sub = parser.add_subparsers(dest="command")

    sub.add_parser("extract", help="Build external link inventory from content files")

    check_parser = sub.add_parser("check", help="Check URLs from inventory")
    check_parser.add_argument("--resume", action="store_true", help="Skip already-checked URLs")
    check_parser.add_argument("--limit", type=int, default=None, help="Max URLs to check this run")

    wayback_parser = sub.add_parser("wayback", help="Look up Wayback Machine snapshots for dead links")
    wayback_parser.add_argument("--resume", action="store_true", help="Skip already-looked-up URLs")
    wayback_parser.add_argument("--limit", type=int, default=None, help="Max URLs to look up this run")

    recheck_parser = sub.add_parser("recheck", help="Re-check 403 URLs with a browser User-Agent")
    recheck_parser.add_argument("--limit", type=int, default=None, help="Max URLs to re-check")

    replace_parser = sub.add_parser("replace", help="Replace dead links with Wayback Machine URLs")
    replace_parser.add_argument("--apply", action="store_true", help="Actually modify files (default is dry run)")

    sub.add_parser("classify", help="Join inventory + results + wayback into the master review sheet")
    sub.add_parser("pages", help="Aggregate master sheet into one row per page (run classify first)")

    args = parser.parse_args()
    if args.command == "extract":
        run_extract()
    elif args.command == "check":
        run_check(resume=args.resume, limit=args.limit)
    elif args.command == "wayback":
        run_wayback(resume=args.resume, limit=args.limit)
    elif args.command == "recheck":
        run_recheck(limit=args.limit)
    elif args.command == "replace":
        run_replace(dry_run=not args.apply)
    elif args.command == "classify":
        run_classify()
    elif args.command == "pages":
        run_pages()
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
